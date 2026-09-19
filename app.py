from io import BytesIO
import os
import re
from docx import Document
from google import genai
from PIL import Image

pypdf_available = True
try:
  from pypdf import PdfReader
except ImportError:
  pypdf_available = False

import gspread
import pandas as pd
import streamlit as st
from google.oauth2.service_account import Credentials

# ---------------------------------------------------------
# KONFIGURASI HALAMAN & KONEKSI GOOGLE SHEETS
# ---------------------------------------------------------
st.set_page_config(
    page_title="Portal Akademik & Asesmen Instansi Pendidikan",
    page_icon="🏛️",
    layout="centered",
    initial_sidebar_state="collapsed",
)


# Fungsi untuk menghubungkan ke Google Sheets via Streamlit Secrets
def get_google_sheet_data():
  try:
    # Mengambil kredensial dari st.secrets
    scope = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive",
    ]
    creds_dict = dict(st.secrets["gcp_service_account"])
    creds = Credentials.from_service_account_info(creds_dict, scopes=scope)
    client = gspread.authorize(creds)

    # Buka Google Sheet berdasarkan Nama File
    sheet_name = st.secrets["sheet_config"]["sheet_name"]
    sheet = client.open(sheet_name).sheet1
    return sheet
  except Exception as e:
    return None


# ---------------------------------------------------------
# CUSTOM CSS: ANIMASI BACKGROUND BERGERAK DINAMIS & WARNA KONTRAST
# ---------------------------------------------------------
st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&display=swap');

    .stApp {
        background: linear-gradient(125deg, #030712 0%, #0f172a 35%, #1e1b4b 70%, #064e3b 100%);
        background-size: 400% 400%;
        animation: gradientAnimation 18s ease infinite;
        font-family: 'Plus Jakarta Sans', sans-serif;
        color: #f8fafc;
    }

    @keyframes gradientAnimation {
        0% { background-position: 0% 50%; }
        50% { background-position: 100% 50%; }
        100% { background-position: 0% 50%; }
    }

    .hero-banner {
        background: linear-gradient(135deg, rgba(30, 58, 138, 0.9) 0%, rgba(15, 23, 42, 0.95) 100%);
        backdrop-filter: blur(12px);
        padding: 32px 24px;
        border-radius: 20px;
        color: #ffffff;
        margin-bottom: 24px;
        border-left: 6px solid #fbbf24;
        border-top: 1px solid rgba(255, 255, 255, 0.2);
    }
    .hero-title {
        font-size: 1.55rem;
        font-weight: 800;
        color: #fbbf24 !important;
        margin: 0;
    }
    .hero-subtitle {
        font-size: 0.86rem;
        color: #cbd5e1 !important;
        margin-top: 8px;
    }

    .dashboard-card {
        background: rgba(15, 23, 42, 0.88);
        backdrop-filter: blur(24px);
        padding: 24px;
        border-radius: 20px;
        border: 1.5px solid rgba(251, 191, 36, 0.35);
        box-shadow: 0 20px 40px -10px rgba(0, 0, 0, 0.8);
        margin-bottom: 20px;
    }

    .stButton>button {
        width: 100%;
        background: linear-gradient(135deg, #d97706 0%, #fbbf24 100%);
        color: #030712 !important;
        font-weight: 800 !important;
        border-radius: 12px !important;
        padding: 0.85rem 1rem !important;
        border: none !important;
    }
    </style>
""",
    unsafe_allow_html=True,
)


# ---------------------------------------------------------
# FUNGSI PEMBENTUK DOKUMEN WORD (.DOCX)
# ---------------------------------------------------------
def buat_file_docx(teks_konten):
  doc = Document()
  teks_bersih = re.sub(
      r"<\s*br\s*/?>", "\n", teks_konten, flags=re.IGNORECASE
  )
  teks_bersih = re.sub(r"</?[a-zA-Z]+[^>]*>", "", teks_bersih)

  lines = teks_bersih.split("\n")
  table_rows = []

  def flush_table():
    nonlocal table_rows
    if table_rows:
      filtered_rows = [
          row
          for row in table_rows
          if not all(
              c in "-:| "
              for c in "".join(
                  [cell for cell in row if isinstance(cell, str)]
              )
          )
      ]
      if filtered_rows:
        num_cols = max(len(row) for row in filtered_rows)
        table = doc.add_table(rows=len(filtered_rows), cols=num_cols)
        table.style = "Table Grid"
        for r_idx, row_data in enumerate(filtered_rows):
          for c_idx, cell_value in enumerate(row_data):
            if c_idx < num_cols:
              table.cell(r_idx, c_idx).text = str(cell_value).strip()
      table_rows = []

  for line in lines:
    stripped = line.strip()
    if stripped.startswith("|") and stripped.endswith("|"):
      cells = [c.strip() for c in stripped.split("|")[1:-1]]
      table_rows.append(cells)
      continue
    else:
      flush_table()

    if not stripped:
      continue

    if stripped.startswith("# "):
      doc.add_heading(stripped.replace("# ", "").strip(), level=1)
    elif stripped.startswith("## "):
      doc.add_heading(stripped.replace("## ", "").strip(), level=2)
    elif stripped.startswith("### "):
      doc.add_heading(stripped.replace("### ", "").strip(), level=3)
    else:
      doc.add_paragraph(stripped)

  flush_table()
  buffer = BytesIO()
  doc.save(buffer)
  buffer.seek(0)
  return buffer


def generate_content_with_retry(
    client, model_name, contents, max_retries=3, delay=3
):
  import time

  for attempt in range(max_retries):
    try:
      return client.models.generate_content(model=model_name, contents=contents)
    except Exception as e:
      if "503" in str(e) or "UNAVAILABLE" in str(e):
        if attempt < max_retries - 1:
          time.sleep(delay * (attempt + 1))
          continue
      raise e


# ---------------------------------------------------------
# INISIALISASI API KEY
# ---------------------------------------------------------
raw_key = st.secrets.get("GEMINI_API_KEY", "") or st.secrets.get(
    "GOOGLE_API_KEY", ""
)
GEMINI_API_KEY = str(raw_key).strip().strip('"').strip("'")

# ---------------------------------------------------------
# NAVIGASI & VALIDASI TOKEN DARI GOOGLE SHEETS
# ---------------------------------------------------------
with st.sidebar:
  st.markdown("### 🏛️ Portal Akademik Resmi")
  st.markdown("---")

  st.markdown("### 🎟️ Aktivasi Token Kuota")
  input_token = st.text_input(
      "Masukkan Token Anda", type="password", placeholder="Contoh: VIP-SITI-10"
  )

  token_aktif = False
  sisa_kuota_sekarang = 0
  row_index = None

  cleaned_token = input_token.strip()

  if cleaned_token != "":
    sheet = get_google_sheet_data()
    if sheet:
      try:
        records = sheet.get_all_records()  # Mengambil seluruh data dari Sheets
        for idx, row in enumerate(records, start=2):  # Baris 2 ke bawah
          if str(row.get("token")).strip() == cleaned_token:
            sisa_kuota_sekarang = int(row.get("kuota", 0))
            row_index = idx
            break

        if row_index is not None:
          if sisa_kuota_sekarang > 0:
            token_aktif = True
            st.success(f"✅ Token Aktif! Sisa Kuota: **{sisa_kuota_sekarang}x**")
          else:
            st.error("❌ Kuota Anda sudah habis di Google Sheets!")
        else:
          st.error("❌ Token tidak ditemukan.")
      except Exception as e:
        st.error(f"Gagal membaca database: {e}")
    else:
      st.error("Gagal terhubung ke Google Sheets.")
  else:
    st.warning("⚠️ Masukkan token untuk mulai menggunakan generator.")

  st.markdown("---")
  menu_pilihan = st.radio(
      "NAVIGASI UTAMA:",
      [
          "📖 Generator Modul Ajar",
          "🎯 Generator CP & ATP",
          "📊 Generator KKTP",
          "📅 Generator Prosem & Prota",
          "📝 Generator Soal Asesmen",
      ],
      index=0,
  )

if not GEMINI_API_KEY:
  st.warning("Mohon masukkan Gemini API Key di panel Secrets.")
  st.stop()

client = genai.Client(api_key=GEMINI_API_KEY)

if "modul_hasil" not in st.session_state:
  st.session_state.modul_hasil = ""

st.markdown(
    """
    <div class="hero-banner">
        <h1 class="hero-title">🏛️ Portal Asisten Akademik & Asesmen</h1>
        <p class="hero-subtitle">Sistem Terintegrasi Berbasis Cloud Sheets.</p>
    </div>
    """,
    unsafe_allow_html=True,
)

if not token_aktif:
  st.markdown('<div class="dashboard-card">', unsafe_allow_html=True)
  st.markdown("### 🔒 Akses Terkunci")
  st.markdown(
      "Masukkan **Token Kuota** yang valid di panel *sidebar* untuk mulai"
      " menggunakan aplikasi."
  )
  st.markdown("</div>", unsafe_allow_html=True)
  st.stop()

# ---------------------------------------------------------
# CONTOH MENU: GENERATOR MODUL AJAR (DENGAN PENGURANGAN KUOTA OTOMATIS)
# ---------------------------------------------------------
if menu_pilihan == "📖 Generator Modul Ajar":
  st.markdown('<div class="dashboard-card">', unsafe_allow_html=True)
  st.markdown("### 📖 Generator Modul Ajar / RPP")

  nama_guru = st.text_input("👤 Nama Guru & Gelar")
  nama_sekolah = st.text_input("🏫 Nama Instansi / Sekolah")
  mapel = st.text_input("📚 Mata Pelajaran")
  topik = st.text_input("💡 Topik / Materi Pokok")

  if st.button("🚀 Buat Modul Ajar Sekarang", type="primary"):
    with st.spinner("Sistem AI sedang merancang Modul Ajar..."):
      try:
        prompt = (
            f"Buatkan Modul Ajar formal untuk Guru {nama_guru} di Sekolah"
            f" {nama_sekolah}, Mapel {mapel}, Topik {topik}."
        )
        response = generate_content_with_retry(
            client, "gemini-3.5-flash", prompt
        )
        st.session_state.modul_hasil = response.text

        # OTOMATIS KURANGI KUOTA DI GOOGLE SHEETS
        new_quota = sisa_kuota_sekarang - 1
        sheet.update_cell(
            row_index, 3, new_quota
        )  # Kolom 3 adalah kolom 'kuota'

        st.success(
            f"Berhasil! Sisa kuota Anda diperbarui di Sheets jadi: {new_quota}x"
        )
      except Exception as e:
        st.error(f"Error: {e}")

  if st.session_state.modul_hasil:
    st.markdown("---")
    st.markdown(st.session_state.modul_hasil)
    file_docx = buat_file_docx(st.session_state.modul_hasil)
    st.download_button(
        "📥 Unduh Modul Ajar (.docx)",
        data=file_docx,
        file_name=f"Modul_Ajar_{mapel}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
  st.markdown("</div>", unsafe_allow_html=True)
